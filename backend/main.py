import io
import json
import os
from datetime import date, timedelta
from typing import Optional

from fastapi import Body, Depends, FastAPI, HTTPException, Query, Response
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from auth import (
    COOKIE_NAME,
    MAX_AGE_SECONDS,
    SITE_PASSWORD,
    make_session_token,
    require_api_key,
    require_session,
)
from db import ensure_indexes, get_db, make_dedupe_key, now_iso

app = FastAPI(title="Job Tracker")

FRONTEND_DIR = os.path.join(os.path.dirname(__file__), "..", "frontend")

STALE_DAYS = 60


@app.on_event("startup")
async def _startup():
    await ensure_indexes()


# ---------------------------------------------------------------- auth -----

class LoginBody(BaseModel):
    password: str


@app.post("/api/login")
async def login(body: LoginBody, response: Response):
    if not SITE_PASSWORD:
        raise HTTPException(500, "Server has no SITE_PASSWORD configured")
    if body.password != SITE_PASSWORD:
        raise HTTPException(401, "Wrong password")
    token = make_session_token()
    response.set_cookie(
        COOKIE_NAME, token, max_age=MAX_AGE_SECONDS, httponly=True,
        samesite="lax", secure=True, path="/",
    )
    return {"ok": True}


@app.post("/api/logout")
async def logout(response: Response):
    response.delete_cookie(COOKIE_NAME, path="/")
    return {"ok": True}


@app.get("/api/me")
async def me(ok: bool = Depends(require_session)):
    return {"ok": True}


# ---------------------------------------------------------------- jobs -----

def _job_out(doc: dict) -> dict:
    doc = dict(doc)
    doc["id"] = str(doc.pop("_id"))
    return doc


@app.get("/api/jobs")
async def list_jobs(
    source: Optional[str] = None,
    applied: Optional[str] = None,
    include_stale: bool = False,
    ok: bool = Depends(require_session),
):
    db = get_db()
    query: dict = {}
    if source:
        query["source"] = source
    if applied in ("yes", "no"):
        query["applied"] = applied == "yes"
    if not include_stale:
        cutoff = (date.today() - timedelta(days=STALE_DAYS)).isoformat()
        # keep it if it's recent OR already applied (never hide applied history)
        query["$or"] = [{"posted_date": {"$gte": cutoff}}, {"applied": True}]
    cursor = db.jobs.find(query).sort([("fit_score", -1), ("posted_date", -1)])
    jobs = [_job_out(d) async for d in cursor]
    return jobs


class AppliedBody(BaseModel):
    applied: bool


@app.patch("/api/jobs/{job_id}")
async def update_job(job_id: str, body: AppliedBody, ok: bool = Depends(require_session)):
    from bson import ObjectId
    db = get_db()
    update = {"applied": body.applied, "updated_at": now_iso()}
    update["applied_at"] = now_iso() if body.applied else None
    res = await db.jobs.update_one({"_id": ObjectId(job_id)}, {"$set": update})
    if res.matched_count == 0:
        raise HTTPException(404, "Job not found")
    return {"ok": True}


@app.get("/api/jobs/{job_id}/resume")
async def get_resume(job_id: str, ok: bool = Depends(require_session)):
    from bson import ObjectId
    db = get_db()
    resume = await db.resumes.find_one({"job_id": job_id})
    if not resume:
        raise HTTPException(404, "No tailored resume for this job yet")
    job = await db.jobs.find_one({"_id": ObjectId(job_id)})
    docx_bytes = _text_to_docx(resume["content_text"])
    filename = "resume_{}.docx".format((job or {}).get("company_or_family", "tailored").replace(" ", "_"))
    return StreamingResponse(
        io.BytesIO(docx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _text_to_docx(text: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            doc.add_paragraph("")
            continue
        is_heading = line.strip() == line.strip().upper() and len(line.strip()) < 60 and any(c.isalpha() for c in line)
        if is_heading:
            p = doc.add_paragraph()
            run = p.add_run(line.strip())
            run.bold = True
            run.font.size = Pt(12)
        else:
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@app.get("/api/summary")
async def summary(ok: bool = Depends(require_session)):
    db = get_db()
    cutoff = (date.today() - timedelta(days=STALE_DAYS)).isoformat()
    active_or_applied = {"$or": [{"posted_date": {"$gte": cutoff}}, {"applied": True}]}
    out = {}
    for source in ("Accenture", "Product"):
        q = {"source": source, **active_or_applied}
        total = await db.jobs.count_documents(q)
        strong = await db.jobs.count_documents({**q, "fit_score": {"$gte": 4}})
        reviewed = await db.jobs.count_documents({**q, "jd_reviewed": True})
        applied_count = await db.jobs.count_documents({"source": source, "applied": True})
        out[source] = {
            "total": total, "strong_or_excellent": strong,
            "jd_reviewed": reviewed, "applied": applied_count,
        }
    out["applied_total"] = await db.jobs.count_documents({"applied": True})
    out["resumes_generated"] = await db.resumes.count_documents({})
    return out


@app.get("/api/runs")
async def list_runs(limit: int = 20, ok: bool = Depends(require_session)):
    db = get_db()
    cursor = db.runs.find({}).sort("run_date", -1).limit(limit)
    out = []
    async for d in cursor:
        d = dict(d)
        d["id"] = str(d.pop("_id"))
        out.append(d)
    return out


# ------------------------------------------------------- internal (pipeline) --

@app.get("/api/internal/config/base_resume")
async def get_base_resume(ok: bool = Depends(require_api_key)):
    db = get_db()
    doc = await db.config.find_one({"_id": "base_resume"})
    if not doc:
        return {"text": "", "updated_at": None}
    return {"text": doc.get("text", ""), "updated_at": doc.get("updated_at")}


class BaseResumeBody(BaseModel):
    text: str
    filename: Optional[str] = None


async def _set_base_resume(text: str, filename: Optional[str]):
    db = get_db()
    await db.config.update_one(
        {"_id": "base_resume"},
        {"$set": {"text": text, "filename": filename, "updated_at": now_iso()}},
        upsert=True,
    )


@app.post("/api/internal/config/base_resume")
async def set_base_resume(body: BaseResumeBody, ok: bool = Depends(require_api_key)):
    await _set_base_resume(body.text, body.filename)
    return {"ok": True}


@app.get("/api/internal/config/base_resume/set")
async def set_base_resume_get(
    text: str = Query(...),
    filename: Optional[str] = Query(default=None),
    ok: bool = Depends(require_api_key),
):
    """GET mirror of the POST above, for a caller that can only make GET
    requests with no custom body (see DAILY_PIPELINE_RUNBOOK.md)."""
    await _set_base_resume(text, filename)
    return {"ok": True}


@app.get("/api/internal/jobs")
async def internal_list_jobs(source: Optional[str] = None, ok: bool = Depends(require_api_key)):
    """Used by the pipeline to see what's already stored, for dedup."""
    db = get_db()
    query = {"source": source} if source else {}
    cursor = db.jobs.find(query, {"dedupe_key": 1, "title": 1, "company_or_family": 1, "posted_date": 1, "apply_url": 1})
    out = []
    async for d in cursor:
        d = dict(d)
        d["id"] = str(d.pop("_id"))
        out.append(d)
    return out


async def _upsert_one_job(job: dict) -> str:
    """Returns 'inserted' or 'updated'."""
    db = get_db()
    key = make_dedupe_key(job)
    job["dedupe_key"] = key
    job.setdefault("applied", False)
    job.setdefault("applied_at", None)
    job.setdefault("resume_tailored", False)
    job["updated_at"] = now_iso()
    existing = await db.jobs.find_one({"dedupe_key": key})
    if existing:
        # never clobber applied state or timestamps set by the user
        job.pop("applied", None)
        job.pop("applied_at", None)
        await db.jobs.update_one({"dedupe_key": key}, {"$set": job})
        return "updated"
    job["created_at"] = now_iso()
    await db.jobs.insert_one(job)
    return "inserted"


@app.post("/api/internal/jobs/upsert")
async def upsert_jobs(jobs: list[dict] = Body(...), ok: bool = Depends(require_api_key)):
    inserted, updated = 0, 0
    for job in jobs:
        result = await _upsert_one_job(job)
        if result == "inserted":
            inserted += 1
        else:
            updated += 1
    return {"inserted": inserted, "updated": updated}


@app.get("/api/internal/jobs/upsert_one")
async def upsert_one_job_get(job: str = Query(...), ok: bool = Depends(require_api_key)):
    """GET mirror of the POST above, one job per call (job=URL-encoded JSON
    object). For a caller that can only make GET requests - see
    DAILY_PIPELINE_RUNBOOK.md."""
    try:
        job_dict = json.loads(job)
    except json.JSONDecodeError as e:
        raise HTTPException(400, f"job must be valid JSON: {e}")
    result = await _upsert_one_job(job_dict)
    return {result: 1}


class ResumeBody(BaseModel):
    job_id: str
    content_text: str
    base_resume_version: Optional[str] = None


async def _save_resume(job_id: str, content_text: str, base_resume_version: Optional[str]):
    from bson import ObjectId
    db = get_db()
    await db.resumes.update_one(
        {"job_id": job_id},
        {"$set": {
            "job_id": job_id,
            "content_text": content_text,
            "base_resume_version": base_resume_version,
            "generated_at": now_iso(),
        }},
        upsert=True,
    )
    await db.jobs.update_one({"_id": ObjectId(job_id)}, {"$set": {"resume_tailored": True}})


@app.post("/api/internal/resumes")
async def save_resume(body: ResumeBody, ok: bool = Depends(require_api_key)):
    await _save_resume(body.job_id, body.content_text, body.base_resume_version)
    return {"ok": True}


@app.get("/api/internal/resumes/save")
async def save_resume_get(
    job_id: str = Query(...),
    content_text: str = Query(...),
    base_resume_version: Optional[str] = Query(default=None),
    ok: bool = Depends(require_api_key),
):
    """GET mirror of the POST above - see DAILY_PIPELINE_RUNBOOK.md."""
    await _save_resume(job_id, content_text, base_resume_version)
    return {"ok": True}


class RunBody(BaseModel):
    run_date: str
    jobs_added: int = 0
    jobs_reviewed: int = 0
    resumes_generated: int = 0
    notes: str = ""


async def _log_run(run_date: str, jobs_added: int, jobs_reviewed: int, resumes_generated: int, notes: str):
    db = get_db()
    doc = {
        "run_date": run_date, "jobs_added": jobs_added, "jobs_reviewed": jobs_reviewed,
        "resumes_generated": resumes_generated, "notes": notes, "finished_at": now_iso(),
    }
    await db.runs.insert_one(doc)


@app.post("/api/internal/runs")
async def log_run(body: RunBody, ok: bool = Depends(require_api_key)):
    await _log_run(body.run_date, body.jobs_added, body.jobs_reviewed, body.resumes_generated, body.notes)
    return {"ok": True}


@app.get("/api/internal/runs/log")
async def log_run_get(
    run_date: str = Query(...),
    jobs_added: int = Query(default=0),
    jobs_reviewed: int = Query(default=0),
    resumes_generated: int = Query(default=0),
    notes: str = Query(default=""),
    ok: bool = Depends(require_api_key),
):
    """GET mirror of the POST above - see DAILY_PIPELINE_RUNBOOK.md."""
    await _log_run(run_date, jobs_added, jobs_reviewed, resumes_generated, notes)
    return {"ok": True}


# ------------------------------------------------------------- static UI ---

app.mount("/static", StaticFiles(directory=FRONTEND_DIR), name="static")


@app.get("/", response_class=HTMLResponse)
async def index():
    with open(os.path.join(FRONTEND_DIR, "index.html"), encoding="utf-8") as f:
        return f.read()


@app.get("/health")
async def health():
    return {"status": "ok"}
